"""Document readers. Each returns a `ParsedDocument`: a title and a list of `Section`s (text + page + heading), so chunks keep
page numbers and section headings. Optional libraries (pdfplumber/pypdf, python-docx, libzim) are imported lazily.

Unsupported or unreadable files raise `SkipFile(reason)`; the builder records the reason in its report.
"""
from __future__ import annotations

import hashlib
import re
from collections.abc import Iterator
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

TEXT_ENCODINGS = ("utf-8", "utf-8-sig", "cp1251", "latin-1")

PLAIN_EXT = {".txt", ".text", ".log", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml", ".xml", ".ini", ".toml", ".cfg",
             ".conf", ".tex", ".srt"}
CODE_EXT = {".py", ".sh", ".bash", ".js", ".ts", ".c", ".h", ".cpp", ".hpp", ".cc", ".java", ".go", ".rs", ".sql", ".rb", ".pl",
            ".lua", ".r", ".m", ".kt", ".swift", ".cs", ".php", ".ps1", ".bat"}
MARKDOWN_EXT = {".md", ".markdown"}
HTML_EXT = {".html", ".htm", ".xhtml"}
SUPPORTED_EXT = PLAIN_EXT | CODE_EXT | MARKDOWN_EXT | HTML_EXT | {".pdf", ".docx", ".zim"}


class SkipFile(Exception):
    """The file cannot be indexed; str(exc) is the reason shown to the user."""


@dataclass
class Section:
    text: str
    page: int | None = None        # 1-based page number (PDF)
    heading: str | None = None     # nearest heading


@dataclass
class ParsedDocument:
    title: str
    kind: str                       # pdf | docx | markdown | text | code | html | zim
    sections: list[Section] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections)


# ---- text decoding -------------------------------------------------------------------------------

def decode_text(data: bytes) -> str:
    """Decode bytes trying utf-8, utf-8-sig, cp1251, latin-1 (in that order); strips a BOM."""
    for enc in TEXT_ENCODINGS:
        try:
            return data.decode(enc).lstrip("﻿")
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")   # unreachable: latin-1 never fails


def looks_binary(data: bytes) -> bool:
    return b"\x00" in data[:8192]


def _norm_ws(s: str) -> str:
    return re.sub(r"[ \t\r\f\v ]+", " ", s).strip()


# ---- plain text / markdown -----------------------------------------------------------------------

_MD_ATX = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_SETEXT = re.compile(r"^\s{0,3}(=+|-+)\s*$")
_MD_FENCE = re.compile(r"^\s{0,3}(```|~~~)")


def parse_markdown(text: str, default_title: str) -> ParsedDocument:
    sections: list[Section] = []
    heading: str | None = None
    title: str | None = None
    buf: list[str] = []
    in_fence = False

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(body, None, heading))
        buf.clear()

    for line in text.splitlines():
        if _MD_FENCE.match(line):
            in_fence = not in_fence
            buf.append(line)
            continue
        if not in_fence:
            m = _MD_ATX.match(line)
            setext = bool(_MD_SETEXT.match(line)) and bool(buf) and bool(buf[-1].strip())
            if m:
                flush()
                heading = m.group(2).strip() or heading
                if title is None and len(m.group(1)) == 1:
                    title = heading
                continue
            if setext:
                h = buf.pop().strip()
                flush()
                heading = h
                if title is None and line.strip().startswith("="):
                    title = h
                continue
        buf.append(line)
    flush()
    return ParsedDocument(title or default_title, "markdown", sections)


def parse_text(text: str, default_title: str, kind: str = "text") -> ParsedDocument:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return ParsedDocument(default_title, kind, [Section(text)] if text else [])


# ---- HTML ----------------------------------------------------------------------------------------

_SKIP_TAGS = {"script", "style", "nav", "footer", "noscript", "iframe", "svg", "template", "form", "button", "select"}
_SKIP_CLASSES = {"mw-editsection", "navbox", "vertical-navbox", "noprint", "mw-jump-link", "catlinks", "printfooter", "toc"}
_VOID = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
_BLOCK = {"p", "div", "li", "ul", "ol", "dl", "dd", "dt", "table", "tr", "section", "article", "blockquote", "pre", "figure",
          "figcaption", "caption", "main", "header", "aside", "details", "summary", "address", "hr", "br", "tbody", "thead"}
_HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple[str, str]] = []     # ("h", text) | ("p", text)
        self.title = ""
        self._buf: list[str] = []
        self._skip_tag: str | None = None
        self._skip_depth = 0
        self._in_title = False
        self._in_heading: str | None = None
        self._hbuf: list[str] = []
        self._pre = 0

    def _flush(self) -> None:
        text = "".join(self._buf)
        self._buf.clear()
        text = text if self._pre else _norm_ws(text)
        if text.strip():
            self.blocks.append(("p", text.strip()))

    def handle_starttag(self, tag: str, attrs) -> None:
        if self._skip_tag is not None:
            if tag == self._skip_tag and tag not in _VOID:
                self._skip_depth += 1
            return
        a = dict(attrs)
        classes = set((a.get("class") or "").split())
        if tag in _SKIP_TAGS or classes & _SKIP_CLASSES or a.get("role") == "navigation":
            if tag not in _VOID:
                self._skip_tag, self._skip_depth = tag, 1
            return
        if tag == "title":
            self._in_title = True
        elif tag in _HEADINGS:
            self._flush()
            self._in_heading, self._hbuf = tag, []
        elif tag in ("td", "th"):
            self._buf.append(" | ")
        elif tag in _BLOCK:
            self._flush()
            if tag == "pre":
                self._pre += 1

    def handle_startendtag(self, tag: str, attrs) -> None:
        if self._skip_tag is None and tag in _BLOCK:
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if self._skip_tag is not None:
            if tag == self._skip_tag:
                self._skip_depth -= 1
                if self._skip_depth == 0:
                    self._skip_tag = None
            return
        if tag == "title":
            self._in_title = False
        elif tag == self._in_heading:
            h = _norm_ws("".join(self._hbuf))
            if h:
                self.blocks.append(("h", h))
            self._in_heading = None
        elif tag in _BLOCK:
            self._flush()
            if tag == "pre":
                self._pre = max(0, self._pre - 1)

    def handle_data(self, data: str) -> None:
        if self._skip_tag is not None:
            return
        if self._in_title:
            self.title += data
        elif self._in_heading:
            self._hbuf.append(data)
        else:
            self._buf.append(data)

    def close(self) -> None:
        super().close()
        self._flush()


def parse_html(html: str, default_title: str, kind: str = "html", title: str | None = None) -> ParsedDocument:
    """HTML → sections split at h1..h6. Removes script/style/nav/footer and similar boilerplate (stdlib parser, no bs4 needed)."""
    p = _HTMLText()
    p.feed(html)
    p.close()
    doc_title = title or _norm_ws(p.title) or next((t for k, t in p.blocks if k == "h"), "") or default_title
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []

    def flush() -> None:
        if buf:
            sections.append(Section("\n\n".join(buf), None, heading))
            buf.clear()

    for k, t in p.blocks:
        if k == "h":
            flush()
            heading = t
        else:
            buf.append(t)
    flush()
    return ParsedDocument(doc_title, kind, sections)


# ---- PDF -----------------------------------------------------------------------------------------

def _clean_pdf_page(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(?=[a-zа-яіїєґ])", r"\1", text)               # de-hyphenate line breaks before lowercase
    paras = re.split(r"\n\s*\n", text)
    out = [_norm_ws(" ".join(p.split("\n"))) for p in paras]
    return "\n\n".join(p for p in out if p)


def read_pdf(path: Path) -> ParsedDocument:
    pages: list[str] = []
    title = ""
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        pdfplumber = None
    if pdfplumber is not None:
        try:
            with pdfplumber.open(path) as pdf:
                title = str((pdf.metadata or {}).get("Title") or "")
                pages = [p.extract_text() or "" for p in pdf.pages]
        except Exception:  # noqa: BLE001 — any parser failure: fall back to pypdf
            pages = []
    if not any(p.strip() for p in pages):
        try:
            from pypdf import PdfReader  # type: ignore
        except ImportError:
            if pdfplumber is None:
                raise SkipFile('PDF support needs pdfplumber or pypdf (pip install "kovcheg[docs]")') from None
            raise SkipFile("no extractable text (scanned PDF? OCR is not supported)") from None
        try:
            r = PdfReader(str(path))
            if r.is_encrypted and not r.decrypt(""):
                raise SkipFile("encrypted PDF")
            title = title or str((r.metadata or {}).get("/Title") or "")
            pages = [p.extract_text() or "" for p in r.pages]
        except SkipFile:
            raise
        except Exception as e:  # noqa: BLE001
            raise SkipFile(f"unreadable PDF: {e}") from None
    sections = [Section(t, i + 1, None) for i, t in enumerate(_clean_pdf_page(p) for p in pages) if t]
    if not sections:
        raise SkipFile("no extractable text (scanned PDF? OCR is not supported)")
    return ParsedDocument(title.strip() or path.stem, "pdf", sections)


# ---- DOCX ----------------------------------------------------------------------------------------

def _docx_heading(par) -> bool:
    name = (getattr(par.style, "name", "") or "").lower()
    return name.startswith(("heading", "title", "заголовок", "назва")) or name in ("subtitle",)


def read_docx(path: Path) -> ParsedDocument:
    try:
        import docx  # type: ignore
        from docx.table import Table  # type: ignore
        from docx.text.paragraph import Paragraph  # type: ignore
    except ImportError:
        raise SkipFile("DOCX support needs python-docx (pip install \"kovcheg[docs]\")") from None
    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise SkipFile(f"unreadable DOCX: {e}") from None
    title = (d.core_properties.title or "").strip()
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []

    def flush() -> None:
        if buf:
            sections.append(Section("\n\n".join(buf), None, heading))
            buf.clear()

    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            par = Paragraph(el, d)
            text = _norm_ws(par.text)
            if not text:
                continue
            if _docx_heading(par):
                flush()
                heading = text
                if not title:
                    title = text
            else:
                buf.append(text)
        elif tag == "tbl":
            rows = []
            for row in Table(el, d).rows:
                cells: list[str] = []
                for c in row.cells:
                    t = _norm_ws(c.text)
                    if t and (not cells or cells[-1] != t):        # merged cells repeat
                        cells.append(t)
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                buf.append("\n".join(rows))
    flush()
    return ParsedDocument(title or path.stem, "docx", sections)


# ---- dispatch ------------------------------------------------------------------------------------

def read_file(path: Path) -> ParsedDocument:
    """Parse one (non-ZIM) file. Raises SkipFile with a reason for unsupported, binary or empty files."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext == ".zim":
        raise SkipFile("ZIM archives are read with iter_zim()")
    if ext and ext not in SUPPORTED_EXT:
        raise SkipFile(f"unsupported file type {ext}")
    data = path.read_bytes()
    if looks_binary(data):
        raise SkipFile("binary content")
    text = decode_text(data)
    if not text.strip():
        raise SkipFile("empty file")
    if ext in MARKDOWN_EXT:
        doc = parse_markdown(text, path.stem)
    elif ext in HTML_EXT:
        doc = parse_html(text, path.stem)
    else:
        doc = parse_text(text, path.stem, "code" if ext in CODE_EXT else "text")
    if not doc.sections:
        raise SkipFile("no text")
    return doc


# ---- ZIM -----------------------------------------------------------------------------------------

@dataclass
class ZimArticle:
    path: str          # entry path inside the archive
    title: str
    sha256: str        # of the raw article bytes
    size: int
    html: bytes

    def parse(self) -> ParsedDocument:
        return parse_html(decode_text(self.html), self.title or self.path, kind="zim", title=self.title or None)


def iter_zim(path: Path) -> Iterator[ZimArticle]:
    """Iterate ALL HTML articles of a ZIM archive (redirects and non-HTML entries are skipped)."""
    try:
        from libzim.reader import Archive  # type: ignore
    except ImportError:
        raise SkipFile("ZIM support needs libzim (pip install \"kovcheg[zim]\")") from None
    try:
        arc = Archive(str(path))
    except Exception as e:  # noqa: BLE001
        raise SkipFile(f"unreadable ZIM: {e}") from None
    for i in range(arc.entry_count):
        try:
            entry = arc._get_entry_by_id(i)
            if entry.is_redirect:
                continue
            item = entry.get_item()
            if "html" not in (item.mimetype or ""):
                continue
            data = bytes(item.content)
        except Exception:  # noqa: BLE001 — one broken entry must not stop a multi-GB import
            continue
        yield ZimArticle(entry.path, entry.title or "", hashlib.sha256(data).hexdigest(), len(data), data)
